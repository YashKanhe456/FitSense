from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = ROOT / "FitSense_Project_Report_Draft.docx"


def set_page_number(paragraph):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def add_footer_with_page_number(section):
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_page_number(paragraph)


def style_document(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)

    for style_name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"

    doc.styles["Title"].font.size = Pt(20)
    doc.styles["Heading 1"].font.size = Pt(16)
    doc.styles["Heading 1"].font.bold = True
    doc.styles["Heading 2"].font.size = Pt(14)
    doc.styles["Heading 2"].font.bold = True


def add_center_text(doc, text, bold=False, size=12):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    return p


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_placeholder_box(doc, title, lines):
    doc.add_paragraph(title, style="Heading 2")
    table = doc.add_table(rows=lines, cols=1)
    table.style = "Table Grid"
    for idx in range(lines):
        cell = table.cell(idx, 0)
        cell.text = "[Placeholder]"
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_code_block(doc, code_text):
    p = doc.add_paragraph()
    run = p.add_run(code_text)
    run.font.name = "Courier New"
    run.font.size = Pt(10)


def add_data_dictionary_table(doc, title, rows):
    doc.add_paragraph(title, style="Heading 2")
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Field Name"
    hdr[1].text = "Data Type"
    hdr[2].text = "Key / Constraint"
    hdr[3].text = "Description"
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value


def main():
    doc = Document()
    style_document(doc)
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.1)
    section.right_margin = Inches(1)
    add_footer_with_page_number(section)

    add_center_text(doc, "A PROJECT REPORT", bold=True, size=16)
    add_center_text(doc, "ON", bold=True, size=14)
    add_center_text(doc, "FITSENSE", bold=True, size=22)
    doc.add_paragraph()
    add_center_text(doc, "Submitted By", bold=True, size=13)
    add_center_text(doc, "Student Name : ________________________________", size=12)
    add_center_text(doc, "Roll Number : _________________________________", size=12)
    add_center_text(doc, "Group Member(s) : _____________________________", size=12)
    doc.add_paragraph()
    add_center_text(doc, "Under the Guidance of", bold=True, size=13)
    add_center_text(doc, "Guide Name : _________________________________", size=12)
    doc.add_paragraph()
    add_center_text(doc, "A project report submitted in partial fulfillment of", size=12)
    add_center_text(doc, "Course : _________________________________", size=12)
    add_center_text(doc, "College : _________________________________", size=12)
    add_center_text(doc, "Academic Year : ___________________________", size=12)

    doc.add_page_break()

    doc.add_paragraph("ACKNOWLEDGEMENT", style="Heading 1")
    doc.add_paragraph(
        "I would like to express my sincere gratitude to my project guide, teachers, department, "
        "and all those who supported the completion of this project report on FitSense. Their "
        "guidance, encouragement, and technical support helped in understanding the design, "
        "implementation, and documentation of the project."
    )
    doc.add_paragraph(
        "I am also thankful to my classmates, friends, and family members for their valuable "
        "feedback and motivation throughout the development of the project. Their suggestions "
        "helped improve the system's functionality, presentation, and documentation quality."
    )
    doc.add_paragraph(
        "Finally, I acknowledge the use of official technical documentation, development tools, "
        "and reference materials related to React, Node.js, Express, PostgreSQL, Vite, and "
        "related libraries that contributed to the successful completion of FitSense."
    )

    doc.add_page_break()

    doc.add_paragraph("INDEX", style="Heading 1")
    index_items = [
        "1. Project Name and Description",
        "2. System Introduction",
        "3. Scope and Limitations of Existing System",
        "4. Proposed System",
        "5. Objectives and Goals of Proposed System",
        "6. Feasibility Study",
        "7. Hardware and Software Specifications",
        "8. E-R Diagram",
        "9. Use Case Diagram",
        "10. Data Dictionary",
        "11. Input Screens and Output Screens",
        "12. Reports",
        "13. Sample Script / Code",
        "14. Future Enhancement",
        "15. References",
    ]
    add_numbered(doc, index_items)

    doc.add_page_break()

    doc.add_paragraph("1. PROJECT NAME AND DESCRIPTION", style="Heading 1")
    doc.add_paragraph("Project Name: FitSense")
    doc.add_paragraph(
        "FitSense is a full-stack e-commerce web application designed to manage the online sale "
        "of fashion products through a responsive storefront and an administrator dashboard. The "
        "system allows customers to browse products, filter items by category and price, add "
        "items to a cart, register or log in, place orders with shipping details, and track their "
        "own order history. At the same time, the administrator can manage products, customers, "
        "orders, and stock through a centralized operational dashboard."
    )
    doc.add_paragraph(
        "The project is developed using React and Vite on the frontend, Node.js and Express on "
        "the backend, and PostgreSQL as the database. The application includes JWT-based "
        "authentication, role-based access control, product image upload support, order status "
        "management, and data-driven analytics for the administrator. FitSense aims to replace "
        "manual or fragmented store operations with a unified digital platform that is easier to "
        "operate, scale, and maintain."
    )

    doc.add_paragraph("2. SYSTEM INTRODUCTION", style="Heading 1")
    doc.add_paragraph(
        "FitSense is a web-based online store management system built to provide a smooth buying "
        "experience for customers and a structured management experience for administrators. "
        "The system gives customers a digital storefront where they can explore featured "
        "products, browse categories such as Shirts, T-Shirts, Jeans, Pants, Formal Wear, and "
        "Casual Wear, view product details, select sizes, manage cart items, and place orders."
    )
    doc.add_paragraph(
        "The administrative side of the system provides tools for adding products, editing product "
        "details, monitoring low-stock items, reviewing recent orders, updating order statuses, "
        "and analyzing sales trends. The application uses a centralized PostgreSQL database, "
        "ensuring that user, product, order, and order item information is stored securely and can "
        "be retrieved quickly when needed."
    )
    doc.add_paragraph(
        "The system improves store operations by combining customer-facing shopping flows and "
        "admin-facing management workflows in one integrated solution. The design is responsive "
        "and intended to work on laptops, desktops, and mobile devices."
    )

    doc.add_paragraph("3. SCOPE AND LIMITATIONS OF EXISTING SYSTEM", style="Heading 1")
    doc.add_paragraph("Scope of Existing System:")
    doc.add_paragraph(
        "Many small retail businesses still depend on social media pages, chat-based order "
        "taking, manual stock records, and disconnected spreadsheets to manage products and "
        "customer orders. In such setups, browsing, ordering, stock management, and reporting "
        "often happen in separate places."
    )
    doc.add_paragraph("Limitations of Existing System:")
    add_bullets(
        doc,
        [
            "Lack of a centralized product catalog for customers.",
            "Manual tracking of orders and stock quantities.",
            "Difficulty in monitoring low-stock items in real time.",
            "No role-based admin panel for controlled product and order management.",
            "Limited analytics for revenue, best-selling products, and customer value.",
            "Higher probability of order handling errors and delayed updates.",
            "Poor customer experience due to scattered information and manual follow-up.",
        ],
    )
    doc.add_paragraph(
        "Because of these limitations, store owners may face inefficiency, inconsistent records, "
        "delayed order processing, and reduced customer satisfaction. Therefore, a structured "
        "online system like FitSense is needed."
    )

    doc.add_paragraph("4. PROPOSED SYSTEM", style="Heading 1")
    doc.add_paragraph(
        "The proposed system, FitSense, is a modern e-commerce application that unifies product "
        "display, user authentication, shopping cart management, order placement, order tracking, "
        "inventory management, and admin analytics in a single platform."
    )
    doc.add_paragraph("Main modules of the proposed system include:")
    add_bullets(
        doc,
        [
            "User registration and login",
            "Product listing and search/filtering",
            "Featured products and related products display",
            "Product details with size selection",
            "Cart and checkout workflow",
            "Order placement and personal order history",
            "Admin dashboard and sales analytics",
            "Product creation, editing, and deletion",
            "Customer management",
            "Order status management",
            "Product image upload support",
        ],
    )
    doc.add_paragraph("Advantages of the proposed system:")
    add_numbered(
        doc,
        [
            "It centralizes products, users, orders, and analytics in one database-driven platform.",
            "It reduces manual stock calculations by updating inventory during order processing.",
            "It improves the shopping experience through searchable and filterable product pages.",
            "It supports role-based access control for secure admin operations.",
            "It provides meaningful business insight through low-stock alerts, revenue summaries, monthly sales trends, top products, and top customers.",
            "It supports responsive usage across multiple devices.",
            "It improves operational speed, accuracy, and transparency.",
        ],
    )
    doc.add_paragraph("Scope of Proposed System:")
    doc.add_paragraph(
        "The scope of FitSense covers customer shopping operations, product and inventory "
        "management, order handling, and store administration. The system is suitable for small "
        "to medium online fashion stores that require a responsive interface, secure user access, "
        "database-backed storage, and a maintainable admin workflow."
    )

    doc.add_paragraph("5. OBJECTIVES AND GOALS OF PROPOSED SYSTEM", style="Heading 1")
    doc.add_paragraph("Objectives:")
    add_numbered(
        doc,
        [
            "To develop a web-based e-commerce platform for product browsing and order placement.",
            "To maintain a centralized PostgreSQL database for products, users, orders, and order items.",
            "To provide secure user authentication using JWT-based access control.",
            "To reduce manual effort in stock tracking and order management.",
            "To support an administrator dashboard for products, customers, orders, and analytics.",
            "To provide responsive access from desktop and mobile devices.",
            "To improve order accuracy, system reliability, and ease of maintenance.",
        ],
    )
    doc.add_paragraph("Goals:")
    add_numbered(
        doc,
        [
            "To create a reliable and user-friendly online shopping experience.",
            "To improve the efficiency of store operations through automation.",
            "To maintain secure and structured data storage.",
            "To provide quick access to order and customer information.",
            "To support future scalability for new features and integrations.",
            "To enable better business decisions using operational analytics.",
        ],
    )

    doc.add_paragraph("6. FEASIBILITY STUDY", style="Heading 1")
    doc.add_paragraph("Technical Feasibility:")
    doc.add_paragraph(
        "The project is technically feasible because it uses widely available and well-documented "
        "technologies such as React, Vite, Node.js, Express, PostgreSQL, JWT, and supporting "
        "libraries. The system can run on standard development hardware and can be deployed on "
        "common hosting platforms."
    )
    doc.add_paragraph("Economic Feasibility:")
    doc.add_paragraph(
        "The system is economically feasible because the major development tools and libraries "
        "used are open-source. It reduces manual administrative effort, lowers the chance of data "
        "errors, and improves operational efficiency, making it cost-effective for small and medium businesses."
    )
    doc.add_paragraph("Operational Feasibility:")
    doc.add_paragraph(
        "FitSense is operationally feasible because it offers separate and intuitive workflows for "
        "customers and administrators. Customers can shop and place orders through a clear UI, "
        "while administrators can manage the store from one dashboard."
    )
    doc.add_paragraph("Schedule Feasibility:")
    doc.add_paragraph(
        "The project is modular in nature and can be developed step by step. Features such as "
        "authentication, product management, cart handling, checkout, and analytics can be built "
        "and tested incrementally within an academic project timeline."
    )
    doc.add_paragraph("Conclusion of Feasibility Study:")
    doc.add_paragraph(
        "Based on the technical, economic, operational, and schedule analysis, FitSense is a "
        "practical and implementable project. It is suitable for academic submission as well as "
        "real-world adaptation for online store management."
    )

    doc.add_paragraph("7. HARDWARE AND SOFTWARE SPECIFICATIONS", style="Heading 1")
    doc.add_paragraph("Hardware Requirements:")
    add_bullets(
        doc,
        [
            "Processor: Intel Core i3 or above",
            "RAM: 4 GB or above",
            "Storage: 20 GB free space or above",
            "Input Devices: Keyboard and Mouse",
            "Internet Connectivity",
            "Smartphone / Laptop / Desktop for testing and usage",
        ],
    )
    doc.add_paragraph("Software Requirements:")
    add_bullets(
        doc,
        [
            "Operating System: Windows 10 / Windows 11 / Linux",
            "Frontend: React, Vite, JavaScript, HTML, CSS",
            "Backend: Node.js, Express.js",
            "Database: PostgreSQL",
            "Authentication: JWT, bcryptjs",
            "API Client / HTTP: Axios",
            "Image Upload Support: Multer and Cloudinary-ready configuration",
            "Browser: Google Chrome, Microsoft Edge, Mozilla Firefox",
            "Development Tools: Visual Studio Code, npm, Git",
        ],
    )

    doc.add_paragraph("8. E-R DIAGRAM", style="Heading 1")
    doc.add_paragraph(
        "The main entities in FitSense are Users, Products, Orders, and Order Items. A user can "
        "place many orders, one order can contain many order items, and each order item refers "
        "to one product."
    )
    add_placeholder_box(doc, "Insert E-R Diagram Here", 8)

    doc.add_paragraph("9. USE CASE DIAGRAM", style="Heading 1")
    doc.add_paragraph("Major actors in the system are Customer and Administrator.")
    add_bullets(
        doc,
        [
            "Customer use cases: Register, Login, Browse Products, View Product Details, Add to Cart, Checkout, View My Orders.",
            "Administrator use cases: Login, Manage Products, Manage Orders, Manage Customers, View Dashboard Analytics, Upload Product Images.",
        ],
    )
    add_placeholder_box(doc, "Insert Use Case Diagram Here", 8)

    doc.add_paragraph("10. DATA DICTIONARY", style="Heading 1")
    add_data_dictionary_table(
        doc,
        "Table: users",
        [
            ("id", "UUID", "Primary Key", "Unique identifier for each user"),
            ("name", "TEXT", "NOT NULL", "Stores full name of user"),
            ("email", "TEXT", "UNIQUE, NOT NULL", "Stores unique login email"),
            ("password", "TEXT", "NOT NULL", "Stores hashed user password"),
            ("role", "TEXT", "Default customer", "Stores role as customer or admin"),
            ("created_at", "TIMESTAMPTZ", "Default NOW()", "User account creation time"),
            ("updated_at", "TIMESTAMPTZ", "Default NOW()", "Last update time"),
        ],
    )
    add_data_dictionary_table(
        doc,
        "Table: products",
        [
            ("id", "UUID", "Primary Key", "Unique product identifier"),
            ("name", "TEXT", "NOT NULL", "Product name"),
            ("description", "TEXT", "NOT NULL", "Detailed product description"),
            ("price", "NUMERIC(10,2)", "CHECK price >= 0", "Product selling price"),
            ("category", "TEXT", "NOT NULL", "Product category"),
            ("image", "TEXT", "NOT NULL", "Image path, asset key, or URL"),
            ("stock", "INTEGER", "Default 0", "Available stock quantity"),
            ("is_featured", "BOOLEAN", "Default FALSE", "Marks featured products"),
            ("sizes", "TEXT[]", "Default {}", "Available size list"),
            ("brand", "TEXT", "Default FitSense", "Brand or store label"),
            ("created_at", "TIMESTAMPTZ", "Default NOW()", "Product creation time"),
            ("updated_at", "TIMESTAMPTZ", "Default NOW()", "Last update time"),
        ],
    )
    add_data_dictionary_table(
        doc,
        "Table: orders",
        [
            ("id", "UUID", "Primary Key", "Unique order identifier"),
            ("user_id", "UUID", "Foreign Key", "References users(id)"),
            ("shipping_full_name", "TEXT", "NOT NULL", "Shipping recipient name"),
            ("shipping_phone", "TEXT", "NOT NULL", "Shipping contact number"),
            ("shipping_address", "TEXT", "NOT NULL", "Shipping address"),
            ("shipping_city", "TEXT", "NOT NULL", "Shipping city"),
            ("shipping_state", "TEXT", "NOT NULL", "Shipping state"),
            ("shipping_pincode", "TEXT", "NOT NULL", "Shipping postal code"),
            ("payment_method", "TEXT", "Default COD", "Selected payment mode"),
            ("items_price", "NUMERIC(10,2)", "Default 0", "Sum of line-item prices"),
            ("shipping_price", "NUMERIC(10,2)", "Default 0", "Shipping charge"),
            ("total_price", "NUMERIC(10,2)", "Default 0", "Grand total price"),
            ("is_paid", "BOOLEAN", "Default FALSE", "Delivery/payment status flag"),
            ("paid_at", "TIMESTAMPTZ", "Nullable", "Payment completion timestamp"),
            ("order_status", "TEXT", "Default Pending", "Pending/Processing/Shipped/Delivered/Cancelled"),
            ("created_at", "TIMESTAMPTZ", "Default NOW()", "Order creation time"),
            ("updated_at", "TIMESTAMPTZ", "Default NOW()", "Last update time"),
        ],
    )
    add_data_dictionary_table(
        doc,
        "Table: order_items",
        [
            ("id", "UUID", "Primary Key", "Unique order item identifier"),
            ("order_id", "UUID", "Foreign Key", "References orders(id)"),
            ("product_id", "UUID", "Foreign Key", "References products(id)"),
            ("name", "TEXT", "NOT NULL", "Stores product name snapshot"),
            ("image", "TEXT", "NOT NULL", "Stores product image snapshot"),
            ("price", "NUMERIC(10,2)", "NOT NULL", "Stores product price at purchase"),
            ("quantity", "INTEGER", "CHECK quantity >= 1", "Purchased quantity"),
            ("size", "TEXT", "Default ''", "Selected product size"),
        ],
    )

    doc.add_paragraph("11. INPUT SCREENS AND OUTPUT SCREENS", style="Heading 1")
    screen_sections = [
        ("Screen 1: Home Page", "Input: No mandatory direct input; category and CTA navigation", "Output: Featured products, category cards, store introduction", "Purpose: Entry point for customers to explore the storefront"),
        ("Screen 2: Register Page", "Input: Name, Email, Password", "Output: User account creation and authentication token", "Purpose: Register a new customer"),
        ("Screen 3: Login Page", "Input: Email, Password", "Output: Successful login or error message", "Purpose: Authenticate existing user"),
        ("Screen 4: Product Listing Page", "Input: Search, Category, Sort, Price Max, In-Stock Filter, Featured Filter", "Output: Filtered product grid", "Purpose: Help customers browse products efficiently"),
        ("Screen 5: Product Details Page", "Input: Size selection, Quantity, Add to Cart", "Output: Product information, related items, cart addition confirmation", "Purpose: Display detailed product information and enable purchase intent"),
        ("Screen 6: Cart Page", "Input: Quantity updates, Item removal", "Output: Updated cart summary and totals", "Purpose: Manage selected products before checkout"),
        ("Screen 7: Checkout Page", "Input: Shipping details and payment method", "Output: Order creation, order total, success redirection", "Purpose: Place customer orders"),
        ("Screen 8: My Orders Page", "Input: Authenticated user session", "Output: List of user orders and statuses", "Purpose: Allow customers to track their orders"),
        ("Screen 9: Admin Login Page", "Input: Admin email and password", "Output: Redirect to admin dashboard", "Purpose: Secure administrator access"),
        ("Screen 10: Admin Dashboard", "Input: Admin session and optional refresh", "Output: Revenue summary, low stock list, recent orders, top products, top customers", "Purpose: Centralized operational control panel"),
        ("Screen 11: Manage Products", "Input: Product form fields and image source", "Output: Product create/update/delete actions", "Purpose: Maintain product catalog"),
        ("Screen 12: Manage Orders", "Input: Order status updates", "Output: Updated order workflow and stock adjustments", "Purpose: Process incoming orders"),
        ("Screen 13: Manage Customers", "Input: Admin view/search actions", "Output: Customer list and activity overview", "Purpose: Review customer accounts"),
    ]
    for title, inp, out, purpose in screen_sections:
        doc.add_paragraph(title, style="Heading 2")
        doc.add_paragraph(inp)
        doc.add_paragraph(out)
        doc.add_paragraph(purpose)
    add_placeholder_box(doc, "Insert Project Screenshots Here", 12)

    doc.add_paragraph("12. REPORTS", style="Heading 1")
    report_items = [
        (
            "Report 1: Sales Summary Report",
            "Shows total revenue, delivered revenue, total orders, pending orders, processing orders, delivered orders, and cancelled orders.",
        ),
        (
            "Report 2: Monthly Sales Trend Report",
            "Displays revenue and order counts month-wise for the last six months.",
        ),
        (
            "Report 3: Category Revenue Breakdown Report",
            "Summarizes category-wise revenue and number of units sold.",
        ),
        (
            "Report 4: Top Products Report",
            "Lists best-selling products based on units sold and revenue generated.",
        ),
        (
            "Report 5: Top Customers Report",
            "Shows customers who generated the highest order value and number of purchases.",
        ),
        (
            "Report 6: Low Stock Alert Report",
            "Highlights products that have reached or crossed the low-stock threshold.",
        ),
        (
            "Report 7: Recent Orders Report",
            "Displays latest order activity with customer details, order items, and total value.",
        ),
        (
            "Report 8: Customer Order History Report",
            "Allows a logged-in customer to view their own placed orders and current status.",
        ),
    ]
    for title, description in report_items:
        doc.add_paragraph(title, style="Heading 2")
        doc.add_paragraph(description)

    doc.add_paragraph("13. SAMPLE SCRIPT / CODE", style="Heading 1")
    doc.add_paragraph("Script 1: PostgreSQL schema initialization")
    add_code_block(
        doc,
        """CREATE TABLE IF NOT EXISTS users (
    id UUID PRIMARY KEY DEFAULT gen_random_uuid(),
    name TEXT NOT NULL,
    email TEXT NOT NULL UNIQUE,
    password TEXT NOT NULL,
    role TEXT NOT NULL DEFAULT 'customer'
);""",
    )
    doc.add_paragraph("Script 2: User login authentication logic")
    add_code_block(
        doc,
        """const userResult = await query(
  'SELECT id, name, email, password, role FROM users WHERE email = $1 LIMIT 1',
  [normalizedEmail]
);
if (user && (await bcrypt.compare(password, user.password))) {
  return { ...safeUser, token: generateToken(safeUser._id) };
}""",
    )
    doc.add_paragraph("Script 3: Order creation logic")
    add_code_block(
        doc,
        """const orderResult = await client.query(
  `INSERT INTO orders (
      user_id, shipping_full_name, shipping_phone, shipping_address,
      shipping_city, shipping_state, shipping_pincode, payment_method,
      items_price, shipping_price, total_price
  ) VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9,$10,$11)
  RETURNING *`
);""",
    )

    doc.add_paragraph("14. FUTURE ENHANCEMENT", style="Heading 1")
    add_numbered(
        doc,
        [
            "Online payment gateway integration beyond cash on delivery.",
            "Wishlist synchronization across authenticated devices.",
            "Advanced coupon and discount management.",
            "Invoice generation and PDF export.",
            "Customer reviews and product ratings.",
            "Email or SMS notifications for order updates.",
            "Analytics export to Excel or PDF.",
            "Mobile application or Progressive Web App support.",
            "Recommendation engine for personalized product suggestions.",
            "Multi-vendor or multi-store support in future versions.",
        ],
    )

    doc.add_paragraph("15. REFERENCES", style="Heading 1")
    references = [
        "React Official Documentation — https://react.dev/",
        "Vite Official Documentation — https://vite.dev/",
        "Node.js Official Documentation — https://nodejs.org/",
        "Express.js Official Documentation — https://expressjs.com/",
        "PostgreSQL Documentation — https://www.postgresql.org/docs/",
        "JWT Introduction and Documentation — https://jwt.io/",
        "Axios Documentation — https://axios-http.com/",
        "Cloudinary Documentation — https://cloudinary.com/documentation",
        "MDN Web Docs — HTML, CSS, JavaScript references — https://developer.mozilla.org/",
        "Official project source code and implementation files of FitSense",
    ]
    add_numbered(doc, references)

    doc.add_page_break()
    doc.add_paragraph("PLACEHOLDERS TO UPDATE BEFORE FINAL SUBMISSION", style="Heading 1")
    add_bullets(
        doc,
        [
            "Student name, roll number, guide name, course, college, and academic year on cover page.",
            "Real project screenshots in the Input / Output Screens section.",
            "Final E-R diagram image.",
            "Final Use Case diagram image.",
            "Any required certificate/declaration page if your college format asks for it.",
        ],
    )

    doc.save(OUTPUT_PATH)
    print(f"Generated: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
